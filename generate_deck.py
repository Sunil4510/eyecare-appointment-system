import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="EyeCare_System_Technical_Review.pptx"):
    prs = Presentation()
    # Set 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_BG_DARK = RGBColor(15, 23, 42)      # Slate 900
    COLOR_CARD_DARK = RGBColor(30, 41, 59)    # Slate 800
    COLOR_PRIMARY = RGBColor(14, 165, 233)    # Sky 500
    COLOR_ACCENT = RGBColor(16, 185, 129)     # Emerald 500
    COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    COLOR_TEXT_LIGHT = RGBColor(226, 232, 240) # Slate 200

    def add_header(slide, title_text, category_text="TECHNICAL ASSESSMENT REVIEW"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.65))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE

    def add_bg(slide, dark=True):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG_DARK if dark else RGBColor(248, 250, 252)
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_DARK):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = RGBColor(51, 65, 85)
        card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.12), Inches(3.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    # Title & Subtitle box
    tbox = s1.shapes.add_textbox(Inches(1.5), Inches(1.7), Inches(10.5), Inches(3.8))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Johnson & Johnson Technology Assessment"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(10)

    p1 = tf.add_paragraph()
    p1.text = "Full-Stack Eye Care Clinic\nAppointment Booking System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_after = Pt(16)

    p2 = tf.add_paragraph()
    p2.text = "Architecture, Design Decisions, Edge-Case Handling & Production Delivery"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Candidate: Sunil Kurapati  |  Role: Full Stack Engineer  |  Stack: React 19 • Express 5 • TypeScript • Docker"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_ACCENT

    # ==========================================
    # SLIDE 2: EXECUTIVE SUMMARY & ASSIGNMENT CONTEXT
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Executive Summary: Objective & Delivery Model")

    # Card 1: The Challenge
    add_card(s2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.1))
    c1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.2), Inches(4.7))
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🎯 Assessment Scope"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(14)
    items1 = [
        "Time-Boxed Target: Deliver within 3-hour envelope.",
        "Full-Stack Coverage: Node.js/Express backend + React 19 SPA frontend.",
        "5 Core User Stories: Auth, catalogue filtering, availability & booking, patient dashboard, optician schedule.",
        "Data Integrity: File-based persistence without an external DB."
    ]
    for it in items1:
        pi = tf1.add_paragraph()
        pi.text = "• " + it
        pi.font.size = Pt(12); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(10)

    # Card 2: Development Methodology
    add_card(s2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.1))
    c2 = s2.shapes.add_textbox(Inches(5.0), Inches(1.8), Inches(3.2), Inches(4.7))
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "⚡ Accelerated Delivery"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(14)
    items2 = [
        "Modern Tooling: Leveraged AI-pair programming to rapidly scaffold, review, and harden.",
        "Engineering Rigor: Verified every line of code; wrote 23 automated tests (Jest + Vitest).",
        "Zero Technical Debt: TypeScript strict mode enabled across both frontend and backend.",
        "Self-Documenting: Generated executive README and complete REST API contract."
    ]
    for it in items2:
        pi = tf2.add_paragraph()
        pi.text = "• " + it
        pi.font.size = Pt(12); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(10)

    # Card 3: Key Deliverables
    add_card(s2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.1))
    c3 = s2.shapes.add_textbox(Inches(9.0), Inches(1.8), Inches(3.3), Inches(4.7))
    tf3 = c3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "🏆 Production Artifacts"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_TEXT_WHITE
    p.space_after = Pt(14)
    items3 = [
        "Repository: Public on GitHub with clean commit history.",
        "Containerization: Complete Docker & Docker Compose setup with Nginx.",
        "Zero-Config Demo: 1-click credential auto-fill for instant reviewer testing.",
        "Flawless Tests: 17 backend + 6 frontend tests passing with 0 warnings."
    ]
    for it in items3:
        pi = tf3.add_paragraph()
        pi.text = "• " + it
        pi.font.size = Pt(12); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(10)

    # ==========================================
    # SLIDE 3: REQUIREMENTS COVERAGE (5 USER STORIES)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Requirements Matrix: 100% User Stories Complete")

    stories = [
        ("US-1: Authentication", "Role-based auth for Patient (James) & Optician (Mary). SHA-256 hashing, credential sanitization from responses, localStorage session persistence.", "POST /login", COLOR_PRIMARY),
        ("US-2: Catalogue Search", "Real-time reactive search across service, clinic, & optician names. 3 multi-criteria dropdowns. Ant Design paginated table.", "GET /catalogue-table", COLOR_ACCENT),
        ("US-3: Slot Booking", "Interactive calendar with past dates disabled. 8 fixed hourly slots (9AM-5PM). Dynamic availability calculation with confirmation step.", "GET & POST /appointments", COLOR_PRIMARY),
        ("US-4: Patient Dashboard", "Chronological list of upcoming appointments. Automatic server-side enrichment of clinic, optician, and service names.", "GET /appointments?patient_id=", COLOR_ACCENT),
        ("US-5: Optician Schedule", "Dynamic header with optician name. Numbered appointment schedule. Interactive Patient Info Modal with complete booking history.", "GET /appointments?optician_id=", COLOR_PRIMARY),
    ]

    for idx, (title, desc, api, col) in enumerate(stories):
        top_y = Inches(1.5 + idx * 1.1)
        add_card(s3, Inches(0.8), top_y, Inches(11.7), Inches(0.95))
        tb = s3.shapes.add_textbox(Inches(1.0), top_y + Inches(0.08), Inches(11.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title}   "
        p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = col
        
        run_api = p.add_run()
        run_api.text = f"[{api}]"
        run_api.font.size = Pt(11); run_api.font.color.rgb = COLOR_TEXT_MUTED

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11); p_desc.font.color.rgb = COLOR_TEXT_LIGHT

    # ==========================================
    # SLIDE 4: THE BIG DIFFERENTIATOR - BUSINESS RULE 4 & EDGE CASES
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Business Rule 4 & Defensive Validation Architecture", "CRITICAL DIFFERENTIATOR")

    # Left Card: Rule 4 Explanation
    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1))
    t1 = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.7))
    tf1 = t1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "⚖️ Business Rule 4: One Clinic Per Day"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    rule_items = [
        "Requirement: 'Each optician can only be at 1 clinic in a day.'",
        "Common Candidate Failure: Most developers only check if a single slot is double-booked, completely missing cross-clinic scheduling.",
        "Our Dual-Layer Defense:",
        "  1. Availability Layer (GET): If Optician is booked at Clinic A on Date X, all 8 slots for Clinic B on Date X are disabled with conflict: true.",
        "  2. Mutation Layer (POST): Backend rejects cross-clinic booking attempts with HTTP 400 Bad Request, even if client validation is bypassed."
    ]
    for it in rule_items:
        pi = tf1.add_paragraph()
        pi.text = it
        pi.font.size = Pt(12); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # Right Card: Other Edge Cases Handled
    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1))
    t2 = s4.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.7))
    tf2 = t2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🛡️ Additional Edge Cases Solved"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(12)

    edge_items = [
        ("No Past Bookings", "Calendar dynamically disables all past dates (disabledDate={(c) => c < today})."),
        ("Strict 8 Hourly Slots", "Restricted strictly to 09:00 AM, 10:00 AM, 11:00 AM, 01:00 PM, 02:00 PM, 03:00 PM, 04:00 PM, 05:00 PM."),
        ("Password Sanitization", "API strips sensitive password hashes before returning user payloads."),
        ("Double-Booking Race Conditions", "Atomic check during appointment creation prevents duplicate slots."),
        ("Async Act Warnings", "Wrapped all React testing renders in waitFor() for zero test console noise.")
    ]
    for title, desc in edge_items:
        pi = tf2.add_paragraph()
        pi.text = f"• {title}: "
        pi.font.bold = True; pi.font.size = Pt(12); pi.font.color.rgb = COLOR_TEXT_WHITE
        run = pi.add_run()
        run.text = desc
        run.font.bold = False; run.font.size = Pt(11); run.font.color.rgb = COLOR_TEXT_MUTED
        pi.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: SYSTEM ARCHITECTURE & DATA FLOW
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Full-Stack System Architecture & Data Enrichment")

    # Layer 1: Frontend
    add_card(s5, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.1))
    tf = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.2), Inches(4.7)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "🖥️ Presentation Layer\n(React 19 + TypeScript)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    for it in [
        "Vite 7.1 with fast HMR",
        "Ant Design 5.28 component system",
        "Tailwind CSS 4 utility styling",
        "React Router 7 client-side SPA routing",
        "AuthContext with localStorage session cache",
        "Atomic component architecture (Molecules & Organisms)"
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # Layer 2: Backend API
    add_card(s5, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.1))
    tf = s5.shapes.add_textbox(Inches(5.0), Inches(1.8), Inches(3.2), Inches(4.7)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "⚙️ Application Layer\n(Node 22 + Express 5)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(12)
    for it in [
        "Express 5 RESTful routing",
        "Server-Side Data Enrichment (resolves clinic, service & optician names)",
        "Rule 4 conflict detection engine",
        "Structured JSON logger (timestamps, metadata)",
        "Global error handling middleware",
        "Strict TypeScript typing throughout"
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # Layer 3: Persistence
    add_card(s5, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.1))
    tf = s5.shapes.add_textbox(Inches(9.0), Inches(1.8), Inches(3.3), Inches(4.7)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "💾 Storage Layer\n(File-Based Fallback)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_TEXT_WHITE
    p.space_after = Pt(12)
    for it in [
        "seed_data/: Immutable initial baseline (users, clinics, opticians, services)",
        "app_data/: Dynamic runtime storage for created appointments",
        "Fallback Mechanism: Reads from app_data/ first, falls back to seed_data/",
        "Git Hygiene: .gitignore + .gitkeep ensures reviewers always start fresh"
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # ==========================================
    # SLIDE 6: TESTING & QUALITY ASSURANCE (23/23 PASSING)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Comprehensive Automated Test Suite: 23 / 23 Passing")

    # Metric Banner
    add_card(s6, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.1), bg_color=RGBColor(6, 78, 59))
    tf = s6.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(0.9)).text_frame
    p = tf.paragraphs[0]
    p.text = "✅ 100% Test Success Rate: 17 Backend Unit/Integration Tests + 6 Frontend Component Tests"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = RGBColor(167, 243, 208)
    p_sub = tf.add_paragraph()
    p_sub.text = "Executed in single-command pipeline (npm test) with 0 failures, 0 regressions, and 0 console warnings."
    p_sub.font.size = Pt(12); p_sub.font.color.rgb = COLOR_TEXT_LIGHT

    # Backend Tests Card
    add_card(s6, Inches(0.8), Inches(2.8), Inches(5.6), Inches(3.9))
    tf = s6.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(5.2), Inches(3.5)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "🧪 Backend Test Suite (Jest & ts-jest)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)
    for it in [
        "health.routes.test.ts: Validates server uptime & presence of all 5 seed datasets.",
        "users.routes.test.ts: Tests patient auth, optician auth, bad passwords, missing payloads.",
        "appointments.routes.test.ts: Validates joined entity responses, sorting, availability calculation.",
        "Rule 4 Isolation Test: Explicitly asserts conflict: true when optician is at another clinic."
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # Frontend Tests Card
    add_card(s6, Inches(6.8), Inches(2.8), Inches(5.7), Inches(3.9))
    tf = s6.shapes.add_textbox(Inches(7.0), Inches(3.0), Inches(5.3), Inches(3.5)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "⚛️ Frontend Test Suite (Vitest & RTL)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(10)
    for it in [
        "Login.test.tsx: Tests input rendering, form submission, and 1-click demo button auto-fill.",
        "Home.test.tsx: Tests appointment table rendering, column presence, and patient schedule.",
        "Catalogue.test.tsx: Tests keyword search bar, 3 filter dropdowns, and service rows.",
        "AppointmentConfirmed.test.tsx: Tests booking confirmation summary card & navigation."
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # ==========================================
    # SLIDE 7: DEVOPS & DOCKER ORCHESTRATION
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "DevOps & Zero-Config Execution: Docker Compose")

    # Command Box
    add_card(s7, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.0), bg_color=RGBColor(30, 58, 138))
    tf = s7.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(0.8)).text_frame
    p = tf.paragraphs[0]
    p.text = "🐳 Single-Command Launch: docker compose up --build"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = RGBColor(191, 219, 254)
    p_sub = tf.add_paragraph()
    p_sub.text = "Reviewers do not even need Node.js installed. Entire full-stack application boots in isolated containers in seconds."
    p_sub.font.size = Pt(12); p_sub.font.color.rgb = COLOR_TEXT_LIGHT

    # Container 1: Backend
    add_card(s7, Inches(0.8), Inches(2.7), Inches(5.6), Inches(4.0))
    tf = s7.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(5.2), Inches(3.6)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "📦 Backend Container (Node 22 Alpine)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)
    for it in [
        "Lightweight node:22-alpine base image.",
        "Compiles TypeScript into dist/ during build.",
        "Bundles seed datasets into dist/seed_data/.",
        "Exposes port 3001.",
        "Built-in Docker healthcheck polling GET /health to ensure readiness before client requests."
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # Container 2: Frontend
    add_card(s7, Inches(6.8), Inches(2.7), Inches(5.7), Inches(4.0))
    tf = s7.shapes.add_textbox(Inches(7.0), Inches(2.9), Inches(5.3), Inches(3.6)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "🌐 Frontend Container (Multi-Stage Nginx)"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(10)
    for it in [
        "Multi-stage build: Node builds React bundle -> Nginx serves static HTML/JS.",
        "Ultra-lightweight (~20MB total image footprint).",
        "Custom nginx.conf with SPA fallback (try_files $uri $uri/ /index.html) prevents 404s on refresh.",
        "Exposes port 3000.",
        "Configured with depends_on: condition: service_healthy for backend sync."
    ]:
        pi = tf.add_paragraph(); pi.text = "• " + it; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_LIGHT; pi.space_after = Pt(8)

    # ==========================================
    # SLIDE 8: CLIENT DEMO & LIVE WALKTHROUGH SCRIPT
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Live Client Walkthrough & Discussion Guide")

    steps = [
        ("1. Show 1-Click Login", "Demonstrate UX focus with quick-fill demo buttons for Patient (James) & Optician (Mary). Highlight SHA-256 sanitization."),
        ("2. Patient Catalogue & Search", "Demonstrate instant multi-filtering (Service, Clinic, Optician) and real-time keyword search across 3 entity types."),
        ("3. Booking Flow & Edge Cases", "Open calendar: show past dates disabled. Select a slot from 8 fixed timeslots. Complete booking and view confirmation card."),
        ("4. Prove Rule 4 Live", "Attempt to book the same Optician at a different clinic on the same date: show real-time conflict banner and slot disabling."),
        ("5. Optician View & History Modal", "Log in as Mary: showcase dynamic schedule header and click a patient to inspect their full medical appointment history."),
        ("6. Quality & DevOps Showcase", "Run 'npm test' in terminal (23/23 passing) and highlight 'docker compose up --build' one-command setup.")
    ]

    for idx, (title, desc) in enumerate(steps):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.5 + row * 1.7)
        add_card(s8, x, y, Inches(5.6), Inches(1.5))
        tf = s8.shapes.add_textbox(x + Inches(0.2), y + Inches(0.12), Inches(5.2), Inches(1.2)).text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY if col == 0 else COLOR_ACCENT
        p.space_after = Pt(4)
        p_desc = tf.add_paragraph(); p_desc.text = desc; p_desc.font.size = Pt(11); p_desc.font.color.rgb = COLOR_TEXT_LIGHT

    # ==========================================
    # SLIDE 9: SUMMARY & WHY HIRE SUNIL
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "Summary: Why This Submission Stands Out")

    pillars = [
        ("🏆 Beyond the Minimum", "Did not just implement basic features; delivered defensive backend validation, data enrichment, and edge-case handling for Rule 4.", COLOR_PRIMARY),
        ("🛡️ Enterprise Rigor", "Zero-warning test suite (23/23 passing), strict TypeScript typings, SHA-256 credential sanitization, and structured JSON logs.", COLOR_ACCENT),
        ("⚡ AI-Augmented Velocity", "Embraced modern AI tools to accelerate scaffolding and test generation, but applied human engineering review to every line.", COLOR_PRIMARY),
        ("📦 Client & Reviewer Empathy", "One-click login buttons, comprehensive documentation, Mermaid diagrams, and Docker Compose orchestration.", COLOR_ACCENT)
    ]

    for idx, (title, desc, col) in enumerate(pillars):
        top_y = Inches(1.6 + idx * 1.25)
        add_card(s9, Inches(0.8), top_y, Inches(11.7), Inches(1.1))
        tf = s9.shapes.add_textbox(Inches(1.0), top_y + Inches(0.12), Inches(11.3), Inches(0.85)).text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = col
        p.space_after = Pt(4)
        p_desc = tf.add_paragraph(); p_desc.text = desc; p_desc.font.size = Pt(12); p_desc.font.color.rgb = COLOR_TEXT_LIGHT

    # Save presentation
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "EyeCare_System_Technical_Review.pptx"
    create_deck(out)
